"""Make a copy of a deck with each slide's speaker notes shown in a panel on the right.

Google Slides imports .pptx speaker notes into its notes pane, which viewers of a
shared link or a PDF never see. This widens every slide by ``--panel-in`` inches
and writes the notes into a text box there, leaving the original content and the
original notes untouched.

    python docs/add_notes_panel.py "docs/GEHA Generator LoRA Finetuning.pptx"
    # -> docs/GEHA Generator LoRA Finetuning (with notes).pptx
"""

from __future__ import annotations

import argparse
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR
from pptx.util import Emu, Inches, Pt

PANEL_FILL = RGBColor(0xEF, 0xEC, 0xE3)
LABEL_COLOR = RGBColor(0x6A, 0x72, 0x80)
TEXT_COLOR = RGBColor(0x16, 0x20, 0x2E)
FONT = "Arial"


def add_panel(slide, left: int, width: int, height: int, notes: str, font_pt: float) -> None:
    panel = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, 0, width, height)
    panel.name = "Speaker notes panel"
    panel.fill.solid()
    panel.fill.fore_color.rgb = PANEL_FILL
    panel.line.fill.background()
    panel.shadow.inherit = False

    margin = Inches(0.4)
    box = slide.shapes.add_textbox(left + margin, margin, width - 2 * margin, height - 2 * margin)
    box.name = "Speaker notes"
    frame = box.text_frame
    frame.word_wrap = True
    frame.vertical_anchor = MSO_ANCHOR.TOP
    for side in ("left", "right", "top", "bottom"):
        setattr(frame, f"margin_{side}", 0)

    label = frame.paragraphs[0].add_run()
    label.text = "SPEAKER NOTES"
    label.font.name, label.font.size, label.font.bold = FONT, Pt(12), True
    label.font.color.rgb = LABEL_COLOR

    for line in (notes.strip() or "(no notes)").splitlines():
        paragraph = frame.add_paragraph()
        paragraph.space_before = Pt(8)
        run = paragraph.add_run()
        run.text = line
        run.font.name, run.font.size = FONT, Pt(font_pt)
        run.font.color.rgb = TEXT_COLOR


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("deck", type=Path)
    parser.add_argument("--output", type=Path)
    parser.add_argument("--panel-in", type=float, default=5.0, help="panel width in inches")
    parser.add_argument("--font-pt", type=float, default=14)
    args = parser.parse_args()

    output = args.output or args.deck.with_name(f"{args.deck.stem} (with notes){args.deck.suffix}")
    deck = Presentation(str(args.deck))
    original_width = deck.slide_width
    panel_width = Inches(args.panel_in)
    deck.slide_width = Emu(original_width + panel_width)

    for slide in deck.slides:
        notes = slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ""
        add_panel(slide, original_width, panel_width, deck.slide_height, notes, args.font_pt)

    deck.save(str(output))
    print(f"Wrote {output} ({len(deck.slides)} slides, {deck.slide_width / 914400:.2f} x {deck.slide_height / 914400:.2f} in)")


if __name__ == "__main__":
    main()
